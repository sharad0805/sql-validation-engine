import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter
from datetime import datetime
import logging

logger = logging.getLogger(__name__)

# ── SEVERITY COLOR MAP ────────────────────────────────────────────────────────
SEVERITY_COLORS = {
    "critical": "FFCCCC",   # red
    "high":     "FFE5CC",   # orange
    "medium":   "FFF9CC",   # yellow
    "low":      "E2EFDA",   # green
}

HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT   = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT    = Font(bold=True, size=13, color="1F4E79")
BORDER_SIDE   = Side(style="thin", color="CCCCCC")
CELL_BORDER   = Border(
    left=BORDER_SIDE, right=BORDER_SIDE,
    top=BORDER_SIDE,  bottom=BORDER_SIDE
)


# ── HELPERS ───────────────────────────────────────────────────────────────────

def _apply_header_row(ws, headers: list[str], row: int = 1):
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.fill   = HEADER_FILL
        cell.font   = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = CELL_BORDER


def _apply_data_row(ws, values: list, row: int, fill_color: str = None):
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col_idx, value=str(value) if value is not None else "NULL")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border    = CELL_BORDER
        if fill_color:
            cell.fill = PatternFill("solid", fgColor=fill_color)


def _auto_column_width(ws, min_width=12, max_width=45):
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=min_width)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(
            max(max_len + 4, min_width), max_width
        )


# ── SHEET 1: SUMMARY ─────────────────────────────────────────────────────────

def _write_summary_sheet(wb: Workbook, violations_df: pd.DataFrame, total_rules: int):
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False

    run_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    total_violations = len(violations_df)
    passed = total_rules - violations_df["rule_id"].nunique()

    # Title block
    ws.merge_cells("A1:F1")
    title_cell = ws["A1"]
    title_cell.value     = "SQL DATA VALIDATION REPORT"
    title_cell.font      = Font(bold=True, size=15, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    ws.merge_cells("A2:F2")
    sub = ws["A2"]
    sub.value     = f"Generated: {run_time}"
    sub.font      = Font(size=10, color="888888")
    sub.alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 20

    ws.row_dimensions[3].height = 10  # spacer

    # KPI boxes
    kpi_data = [
        ("Total Rules",      total_rules,      "1F4E79", "FFFFFF"),
        ("Rules Passed",     passed,            "375623", "FFFFFF"),
        ("Total Violations", total_violations,  "C00000", "FFFFFF"),
        ("Critical",         len(violations_df[violations_df["severity"] == "critical"]), "C00000", "FFFFFF"),
        ("High",             len(violations_df[violations_df["severity"] == "high"]),     "C55A11", "FFFFFF"),
        ("Medium",           len(violations_df[violations_df["severity"] == "medium"]),   "7F6000", "FFFFFF"),
    ]

    for col_idx, (label, value, bg, fg) in enumerate(kpi_data, start=1):
        label_cell = ws.cell(row=4, column=col_idx, value=label)
        value_cell = ws.cell(row=5, column=col_idx, value=value)
        for cell in [label_cell, value_cell]:
            cell.fill      = PatternFill("solid", fgColor=bg)
            cell.font      = Font(color=fg, bold=True, size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = CELL_BORDER
        ws.row_dimensions[4].height = 22
        ws.row_dimensions[5].height = 28

    ws.row_dimensions[6].height = 15  # spacer

    # Violations by severity table
    severity_title = ws.cell(row=7, column=1, value="Violations by Severity")
    severity_title.font = TITLE_FONT
    ws.row_dimensions[7].height = 22

    _apply_header_row(ws, ["Severity", "Count", "Rules Affected"], row=8)
    ws.row_dimensions[8].height = 20

    row = 9
    for sev in ["critical", "high", "medium", "low"]:
        subset = violations_df[violations_df["severity"] == sev]
        if subset.empty:
            continue
        color = SEVERITY_COLORS[sev]
        _apply_data_row(ws, [sev.upper(), len(subset), subset["rule_id"].nunique()], row=row, fill_color=color)
        ws.row_dimensions[row].height = 18
        row += 1

    ws.row_dimensions[row].height = 15  # spacer
    row += 1

    # Violations by rule table
    rule_title = ws.cell(row=row, column=1, value="Violations by Rule")
    rule_title.font = TITLE_FONT
    ws.row_dimensions[row].height = 22
    row += 1

    _apply_header_row(ws, ["Rule ID", "Description", "Table", "Column", "Severity", "Count"], row=row)
    ws.row_dimensions[row].height = 20
    row += 1

    for rule_id, group in violations_df.groupby("rule_id"):
        first = group.iloc[0]
        color = SEVERITY_COLORS.get(first["severity"], "FFFFFF")
        _apply_data_row(ws, [
            rule_id, first["description"], first["table"],
            first["column"], first["severity"].upper(), len(group)
        ], row=row, fill_color=color)
        ws.row_dimensions[row].height = 18
        row += 1

    _auto_column_width(ws)


# ── SHEET 2: ALL VIOLATIONS ───────────────────────────────────────────────────

def _write_violations_sheet(wb: Workbook, violations_df: pd.DataFrame):
    ws = wb.create_sheet(title="All Violations")
    ws.sheet_view.showGridLines = False

    headers = ["Rule ID", "Table", "Column", "Rule Type",
               "Severity", "Row ID", "Violated Value", "Description"]
    _apply_header_row(ws, headers, row=1)
    ws.row_dimensions[1].height = 22

    for row_idx, (_, row_data) in enumerate(violations_df.iterrows(), start=2):
        color = SEVERITY_COLORS.get(row_data["severity"], "FFFFFF")
        _apply_data_row(ws, [
            row_data["rule_id"],
            row_data["table"],
            row_data["column"],
            row_data["rule_type"],
            row_data["severity"].upper(),
            row_data["row_id"],
            row_data["violated_value"],
            row_data["description"],
        ], row=row_idx, fill_color=color)
        ws.row_dimensions[row_idx].height = 18

    _auto_column_width(ws)


# ── MAIN GENERATOR ────────────────────────────────────────────────────────────

def generate_report(
    violations_df: pd.DataFrame,
    total_rules: int,
    output_dir: str = "reports"
) -> str:

    os.makedirs(output_dir, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = os.path.join(output_dir, f"validation_report_{timestamp}.xlsx")

    wb = Workbook()

    if violations_df.empty:
        ws = wb.active
        ws.title = "Summary"
        ws["A1"] = "✅ All validation rules passed. No violations found."
        ws["A1"].font = Font(bold=True, size=13, color="375623")
    else:
        _write_summary_sheet(wb, violations_df, total_rules)
        _write_violations_sheet(wb, violations_df)

    wb.save(output_path)
    logger.info(f"Report saved: {output_path}")
    return output_path